"""Read-only Dataset B inventory. Does not download or build Dataset B."""
from pathlib import Path
import json
import openpyxl

BASE=Path(__file__).resolve().parents[1]
ROOT=BASE.parent

def main():
    p=ROOT/'Serrano_SIGG2021_PerceptualRatings.xlsx'
    w=openpyxl.load_workbook(p,read_only=True,data_only=True)
    sheets={s.title:{'rows_excluding_header':s.max_row-1,'columns':s.max_column,'headers':list(next(s.values))} for s in w}
    median=w['MedianRating_per_image']
    image_names={r[0] for r in median.iter_rows(min_row=2,values_only=True) if r[0]}
    files=list((ROOT/'material-illumination-geometry-main/test').rglob('*.jpg'))
    matched=[p.relative_to(ROOT).as_posix() for p in files if p.name in image_names]
    result={'status':'inventory_only_awaiting_user_decision','source_workbook':p.name,'sheets':sheets,'median_rating_unique_images':len(image_names),'local_test_jpg_count':len(files),'matched_local_test_images':matched,'missing_images_relative_to_test_folder':len(image_names)-len(matched),'download_attempted':False,'cleaning_performed':False,'note':'all_data row count is workbook content, not a claim of author-filtered valid responses. No training or Dataset B conversion performed.'}
    (BASE/'metadata/material_inventory.json').write_text(json.dumps(result,indent=2,ensure_ascii=False),encoding='utf-8')
    print(json.dumps(result,indent=2,ensure_ascii=False))

if __name__=='__main__': main()
